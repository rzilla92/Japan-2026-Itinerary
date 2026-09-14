import json
import re
from datetime import date, datetime, time, timedelta
from pathlib import Path

import openpyxl

SOURCE = Path("source/Japan 2026 - Day itinerary (Organised).xlsx")
OUT = Path("js/data.js")


def value(v):
    if isinstance(v, (datetime, date)):
        return v.isoformat()
    if isinstance(v, time):
        return v.strftime("%H:%M")
    return v


def row_values(sheet, row):
    return [value(sheet.cell(row, col).value) for col in range(1, sheet.max_column + 1)]


def filled_rows(sheet):
    return [[r, row_values(sheet, r)] for r in range(1, sheet.max_row + 1) if any(v is not None for v in row_values(sheet, r))]


book = openpyxl.load_workbook(SOURCE, data_only=True)
days = []
for sheet in book.worksheets:
    if not re.match(r"\d{2}-\d{2} Day", sheet.title):
        continue
    rows = filled_rows(sheet)
    meta = rows[1][1] if len(rows) > 1 else []
    sections = {"places": [], "itinerary": [], "recommendations": []}
    mode = None
    category = None
    for r, cells in rows[2:]:
        first = cells[0] if len(cells) else None
        if first == "Main Places":
            mode = "places"
            continue
        if first == "Time Range":
            mode = "itinerary"
            continue
        if first == "Recommendations & Options":
            mode = "recommendations"
            continue
        if mode == "places" and first:
            sections["places"].append({"name": first, "description": cells[1] if len(cells) > 1 else None})
        elif mode == "itinerary" and first:
            sections["itinerary"].append({"time": first, "activity": cells[1] if len(cells) > 1 else None, "details": cells[2] if len(cells) > 2 else None, "extra": [v for v in cells[3:] if v is not None]})
        elif mode == "recommendations" and first:
            if all(v is None for v in cells[1:]):
                category = first
                sections["recommendations"].append({"category": category, "items": []})
            else:
                if not sections["recommendations"]:
                    category = "Recommendations"
                    sections["recommendations"].append({"category": category, "items": []})
                sections["recommendations"][-1]["items"].append({"label": first, "name": cells[1] if len(cells)>1 else None, "details": cells[2] if len(cells)>2 else None, "extra": [v for v in cells[3:] if v is not None]})
    days.append({"sheet": sheet.title, "day": meta[0], "date": meta[1], "location": meta[2], **sections})

transport_sheet = book["Transport Summary"]
transport = []
current_day = None
for r, cells in filled_rows(transport_sheet)[1:]:
    if cells[0]:
        current_day = cells[0]
    if len(cells) > 1 and cells[1] == "EN":
        transport.append({"day": current_day, "english": {"from": cells[2], "to": cells[3], "payment": cells[4], "cost": cells[5], "mode": cells[6], "departure": cells[7], "arrival": cells[8], "duration": cells[9], "line": cells[10], "notes": cells[11]}, "japanese": None, "alternative": None})
    elif len(cells) > 1 and cells[1] == "JP" and transport:
        transport[-1]["japanese"] = {"from": cells[2], "to": cells[3], "payment": cells[4], "cost": cells[5], "mode": cells[6], "departure": cells[7], "arrival": cells[8], "duration": cells[9], "line": cells[10], "notes": cells[11]}
        if (cells[7], cells[8]) != (transport[-1]["english"]["departure"], transport[-1]["english"]["arrival"]) and (cells[7] or cells[8]):
            transport[-1]["alternative"] = {"departure": cells[7], "arrival": cells[8], "duration": cells[9]}
    elif current_day and len(cells) > 2 and cells[2] and transport and cells[1] is None and any('\u3040' <= ch <= '\u9fff' for ch in str(cells[2])):
        transport[-1]["japanese"] = {"from": cells[2], "to": cells[3], "payment": cells[4], "cost": cells[5], "mode": cells[6], "departure": cells[7], "arrival": cells[8], "duration": cells[9], "line": cells[10], "notes": cells[11]}
    elif current_day and len(cells) > 2 and cells[2] and transport and cells[1] is None:
        transport.append({"day": current_day, "english": {"from": cells[2], "to": cells[3], "payment": cells[4], "cost": cells[5], "mode": cells[6], "departure": cells[7], "arrival": cells[8], "duration": cells[9], "line": cells[10], "notes": cells[11]}, "japanese": {"from": cells[2] if any('\u3040' <= ch <= '\u9fff' for ch in str(cells[2])) else None, "to": cells[3] if any('\u3040' <= ch <= '\u9fff' for ch in str(cells[3])) else None}})

acc_sheet = book["Accomodations"]
accommodations = []
for r in range(2, 16, 2):
    en, jp = row_values(acc_sheet, r), row_values(acc_sheet, r + 1)
    accommodations.append({"date": en[0], "city": en[1], "hotel": en[2], "nights": en[3], "address": en[4], "phone": en[5], "checkInOut": en[6], "laundry": en[7], "japanese": {"date": jp[0], "city": jp[1], "hotel": jp[2], "nights": jp[3], "address": jp[4], "phone": jp[5], "checkInOut": jp[6], "laundry": jp[7]}})

for day in days:
    try:
        day_date = datetime.strptime(day["date"], "%d-%m-%y").date()
    except (TypeError, ValueError):
        day_date = None
    day["base"] = None
    if day_date:
        for stay in accommodations:
            check_in = datetime.fromisoformat(stay["date"]).date()
            if check_in <= day_date < check_in + timedelta(days=int(stay["nights"])):
                day["base"] = stay["city"]
                break

luggage_sheet = book["Luggage Forwarding"]
luggage_destinations = []
for start in (1, 7):
    rows = [row_values(luggage_sheet, r) for r in range(start, start + 5)]
    luggage_destinations.append({"city": rows[0][0], "cityJapanese": rows[1][0], "date": rows[3][0], "hotel": rows[0][2], "hotelJapanese": rows[0][5], "recipient": rows[1][2], "recipientJapanese": rows[1][5], "checkIn": rows[2][2], "checkInJapanese": rows[2][5], "phone": rows[3][2], "phoneJapanese": rows[3][5], "address": rows[4][2], "addressJapanese": rows[4][5]})
luggage_phrases = []
for r in range(15, 31):
    cells = row_values(luggage_sheet, r)
    if cells[0]: luggage_phrases.append({"english": cells[0], "japanese": cells[1], "romanization": cells[2]})
luggage_fields = []
for r in range(15, 20):
    cells = row_values(luggage_sheet, r)
    luggage_fields.append({"field": cells[4], "japanese": cells[5], "example": cells[6]})

phrase_sheet = book["Useful Phrases"]
phrases = []
for r, cells in filled_rows(phrase_sheet)[1:]:
    if cells[0]: phrases.append({"category": cells[0], "english": cells[1], "japanese": cells[2], "romanization": cells[3], "context": cells[4] if len(cells) > 4 else None})

timeline = []
timeline_sheet = book["Timeline"]
for col in range(1, 18):
    timeline.append({"day": timeline_sheet.cell(60, col).value, "date": value(timeline_sheet.cell(61, col).value), "location": timeline_sheet.cell(62, col).value, "detail": timeline_sheet.cell(63, col).value})

gantt_blocks = []
for row in range(62, 78):
    for col in range(1, 18):
        cell = timeline_sheet.cell(row, col)
        if cell.value is None:
            continue
        merge = next((m for m in timeline_sheet.merged_cells.ranges if m.min_row == row and m.min_col == col), None)
        gantt_blocks.append({"row": row, "start": col - 1, "end": (merge.max_col if merge else col) - 1, "name": cell.value})

transport_summary = {"totalPerPerson": transport_sheet["F68"].value, "breakdown": [{"method": transport_sheet["E70"].value, "cost": transport_sheet["F70"].value}, {"method": transport_sheet["E71"].value, "cost": transport_sheet["F71"].value}, {"method": transport_sheet["E72"].value, "cost": transport_sheet["F72"].value}]}
journey_endpoints = [
    {"day": timeline_sheet["A1"].value, "date": value(timeline_sheet["A2"].value), "location": f"{timeline_sheet['A3'].value} → {timeline_sheet['A4'].value}", "base": timeline_sheet["A4"].value},
    {"day": timeline_sheet["Q1"].value, "date": value(timeline_sheet["Q2"].value), "location": f"{timeline_sheet['P18'].value} → {timeline_sheet['Q19'].value}", "base": timeline_sheet["Q19"].value},
]
model = {"days": days, "timeline": timeline, "ganttBlocks": gantt_blocks, "journeyEndpoints": journey_endpoints, "transport": transport, "transportSummary": transport_summary, "accommodations": accommodations, "luggage": {"destinations": luggage_destinations, "phrases": luggage_phrases, "fields": luggage_fields}, "phrases": phrases}
OUT.parent.mkdir(exist_ok=True)
OUT.write_text("/* Generated from the source workbook. Do not edit manually. */\nconst TRIP_DATA = " + json.dumps(model, ensure_ascii=False, indent=2) + ";\n", encoding="utf-8")
